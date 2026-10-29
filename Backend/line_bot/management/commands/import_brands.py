"""
Import Brand_Database from Excel (.xlsx) or CSV into Django database.

Excel:
    python manage.py import_brands financial_data.xlsx --dry-run
    python manage.py import_brands financial_data.xlsx

CSV:
    python manage.py import_brands brand_database.csv --dry-run
    python manage.py import_brands brand_database.csv

Notes:
- Excel must contain a sheet named "Brand_Database".
- Required columns: Keyword, Category, Type.
- Only global categories (user=None) are allowed.
- Existing brands are updated; no brands are deleted.
- Invalid data aborts the entire import.
"""

import csv
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from categories.models import Brand, Category


class Command(BaseCommand):
    help = (
        "Import global brand keywords from Excel or CSV "
        "without deleting existing records."
    )

    def add_arguments(self, parser):
        parser.add_argument(
            "source_file",
            help="Path to .xlsx or .csv file",
        )

        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Validate data without saving changes to the database",
        )

    def read_excel(self, path):
        """Read Brand_Database worksheet from Excel."""

        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise CommandError(
                "ยังไม่ได้ติดตั้ง openpyxl กรุณารัน: "
                "python -m pip install openpyxl"
            ) from exc

        try:
            workbook = load_workbook(
                filename=path,
                read_only=True,
                data_only=True,
            )
        except (OSError, ValueError) as exc:
            raise CommandError(
                f"ไม่สามารถเปิดไฟล์ Excel ได้: {exc}"
            ) from exc

        try:
            if "Brand_Database" not in workbook.sheetnames:
                raise CommandError(
                    "ไม่พบชีต 'Brand_Database' ในไฟล์ Excel"
                )

            sheet = workbook["Brand_Database"]
            data = sheet.iter_rows(values_only=True)

            try:
                header_row = next(data)
            except StopIteration:
                raise CommandError(
                    "ชีต Brand_Database ไม่มีข้อมูล"
                )

            headers = [
                str(value).strip() if value is not None else ""
                for value in header_row
            ]

            rows = []

            for row in data:
                rows.append(dict(zip(headers, row)))

            return rows

        finally:
            workbook.close()

    def read_csv(self, path):
        """Read Brand_Database from UTF-8 CSV."""

        try:
            with path.open(
                "r",
                encoding="utf-8-sig",
                newline="",
            ) as file:
                return list(csv.DictReader(file))

        except (OSError, UnicodeError, csv.Error) as exc:
            raise CommandError(
                f"ไม่สามารถอ่านไฟล์ CSV ได้: {exc}"
            ) from exc

    def handle(self, *args, **options):
        path = Path(options["source_file"])
        dry_run = options["dry_run"]

        # ------------------------------------
        # 1. Check source file
        # ------------------------------------

        if not path.is_file():
            raise CommandError(
                f"ไม่พบไฟล์: {path}"
            )

        extension = path.suffix.lower()

        if extension == ".xlsx":
            rows = self.read_excel(path)

        elif extension == ".csv":
            rows = self.read_csv(path)

        else:
            raise CommandError(
                "รองรับเฉพาะไฟล์ .xlsx และ .csv"
            )

        if not rows:
            raise CommandError(
                "ไม่พบข้อมูลแบรนด์ในไฟล์"
            )

        # ------------------------------------
        # 2. Validate required columns
        # ------------------------------------

        required_columns = {
            "Keyword",
            "Category",
            "Type",
        }

        missing_columns = required_columns - set(rows[0].keys())

        if missing_columns:
            raise CommandError(
                "ไม่พบคอลัมน์ที่จำเป็น: "
                + ", ".join(sorted(missing_columns))
            )

        # ------------------------------------
        # 3. Validate brand data
        # ------------------------------------

        invalid = []
        prepared = {}

        for row_no, row in enumerate(rows, start=2):

            keyword = str(
                row.get("Keyword") or ""
            ).strip().casefold()

            category_name = str(
                row.get("Category") or ""
            ).strip()

            transaction_type = str(
                row.get("Type") or ""
            ).strip().lower()

            # Skip empty rows
            if (
                not keyword
                and not category_name
                and not transaction_type
            ):
                continue

            # Validate keyword
            if not keyword or len(keyword) > 150:
                invalid.append(
                    f"Row {row_no}: invalid Keyword"
                )
                continue

            # Validate transaction type
            if transaction_type not in (
                "income",
                "expense",
            ):
                invalid.append(
                    f"Row {row_no}: invalid Type "
                    f"{transaction_type!r}"
                )
                continue

            # Validate category
            if not category_name:
                invalid.append(
                    f"Row {row_no}: missing Category"
                )
                continue

            category = Category.objects.filter(
                user__isnull=True,
                name=category_name,
                transaction_type=transaction_type,
            ).first()

            if category is None:
                invalid.append(
                    f"Row {row_no}: "
                    f"missing global category "
                    f"{category_name!r} "
                    f"({transaction_type})"
                )
                continue

            key = (
                keyword,
                transaction_type,
            )

            # Prevent conflicting duplicate keywords
            if key in prepared:
                if prepared[key] != category.pk:
                    invalid.append(
                        f"Row {row_no}: "
                        f"conflicting duplicate Keyword "
                        f"{keyword!r}"
                    )
                continue

            prepared[key] = category.pk

        # ------------------------------------
        # 4. Abort on invalid data
        # ------------------------------------

        if invalid:
            error_message = (
                f"พบข้อมูลไม่ถูกต้อง {len(invalid)} รายการ\n"
                "Import aborted; no changes made:\n"
                + "\n".join(invalid[:40])
            )

            if len(invalid) > 40:
                error_message += (
                    f"\n... และอีก {len(invalid) - 40} รายการ"
                )

            raise CommandError(error_message)

        if not prepared:
            raise CommandError(
                "ไม่พบข้อมูลแบรนด์ที่สามารถนำเข้าได้"
            )

        # ------------------------------------
        # 5. Dry run
        # ------------------------------------

        if dry_run:
            self.stdout.write(
                self.style.SUCCESS(
                    f"DRY RUN OK: "
                    f"{len(prepared)} brand keywords valid. "
                    "No data changed."
                )
            )
            return

        # ------------------------------------
        # 6. Import into database
        # ------------------------------------

        created = 0
        updated = 0

        with transaction.atomic():

            for (
                keyword,
                transaction_type,
            ), category_id in prepared.items():

                _, was_created = Brand.objects.update_or_create(
                    keyword=keyword,
                    transaction_type=transaction_type,
                    defaults={
                        "category_id": category_id,
                    },
                )

                if was_created:
                    created += 1
                else:
                    updated += 1

        # ------------------------------------
        # 7. Success message
        # ------------------------------------

        self.stdout.write(
            self.style.SUCCESS(
                f"Import completed: "
                f"{created} new, "
                f"{updated} updated brand keywords."
            )
        )